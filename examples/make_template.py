"""Build examples/acme-template.docx: a made-up consulting firm's house template."""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0F, 0x76, 0x6E)


def main(path: str) -> None:
    doc = Document()
    for level, size in ((1, 20), (2, 14), (3, 12)):
        h = doc.styles[f"Heading {level}"]
        h.font.name, h.font.size, h.font.color.rgb = "Georgia", Pt(size), NAVY
    body = doc.styles.add_style("ACME Body", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name, body.font.size = "Georgia", Pt(10.5)
    body.paragraph_format.space_after = Pt(8)
    cell = doc.styles.add_style("ACME Table Text", WD_STYLE_TYPE.PARAGRAPH)
    cell.font.name, cell.font.size = "Arial", Pt(9)
    caption = doc.styles["Caption"]
    caption.font.color.rgb = TEAL

    # Cover page
    for _ in range(6):
        doc.add_paragraph()
    name = doc.add_paragraph("ACME ENGINEERING")
    name.runs[0].font.size, name.runs[0].font.bold, name.runs[0].font.color.rgb = Pt(28), True, NAVY
    doc.add_paragraph("[Report title]").runs[0].font.size = Pt(20)
    doc.add_paragraph("Prepared for [Client]")
    doc.add_paragraph("[Date]")
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)

    body_section.header.is_linked_to_previous = False
    header = body_section.header.paragraphs[0]
    header.text = "ACME Engineering  |  Confidential"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    body_section.footer.is_linked_to_previous = False
    body_section.footer.paragraphs[0].text = "acme-engineering.example  ·  Registered in England No. 0000000"

    # Example body so `rework learn` can see the house conventions in use.
    doc.add_heading("Introduction", level=1)
    for _ in range(3):
        doc.add_paragraph("Body text in the ACME house style. " * 6, style="ACME Body")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Medium Shading 1 Accent 1"
    for row in table.rows:
        for c in row.cells:
            c.paragraphs[0].text = "Value"
            c.paragraphs[0].style = cell
    doc.add_paragraph("Table 1: Example", style="Caption")
    doc.save(path)


if __name__ == "__main__":
    main(os.path.join(os.path.dirname(__file__), "acme-template.docx"))
