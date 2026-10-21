import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor
from lxml import etree

from rework.cli import build, main
from rework.learn import learn
from rework.ooxml import R_NS, W_NS, resolve_styles, style_names, w

RT_HYPERLINK = RELATIONSHIP_TYPE.HYPERLINK

pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="needs pandoc")

MARKDOWN = """---
title: Bridge Deck Assessment
author: Example Consulting
date: October 2026
---

# Introduction

This report summarises the condition of the deck.

- first point
- second point

| Core | Chloride |
|------|----------|
| C1   | 0.12     |

Table: Chloride by core
"""


def make_template(path):
    """A small house template: cover section, a custom body style, a table style, a header."""
    doc = Document()
    body = doc.styles.add_style("House Body", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name = "Georgia"
    body.font.size = Pt(11)
    doc.add_paragraph("ACME Engineering", style="Title")
    doc.add_paragraph("Report title goes here")
    doc.add_paragraph("Prepared for: Client name")
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "ACME Confidential"
    doc.add_heading("Background", level=1)
    for _ in range(3):
        doc.add_paragraph("House body text " * 10, style="House Body")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Light Grid Accent 1"
    table.cell(0, 0).text = "a"
    doc.save(path)


def read(path, part):
    return etree.fromstring(zipfile.ZipFile(path).read(part))


@pytest.fixture
def files(tmp_path):
    tpl, md = tmp_path / "house.docx", tmp_path / "report.md"
    make_template(tpl)
    md.write_text(MARKDOWN)
    return tpl, md, tmp_path / "out.docx"


def test_learn_finds_house_style(files):
    tpl, _, _ = files
    style = learn(str(tpl))
    assert style["body"] == "House Body"  # names, not ids
    assert style["table"] == "Light Grid Accent 1"
    assert style["cover"] == 4  # three lines plus the section break
    assert style["body_section"] == 1
    assert "Report title goes here" in style["cover_lines"]


def test_build_applies_house_style(files):
    tpl, md, out = files
    style = learn(str(tpl))
    style["replace"] = [["Report title goes here", "{title}"], ["Client name", "{client}"]]
    missing = build(str(md), str(tpl), str(out), style, {"client": "City of Springfield"})
    assert missing == []

    doc = read(out, "word/document.xml")
    texts = ["".join(t.text or "" for t in p.iter(w("t"))) for p in doc.iter(w("p"))]
    assert texts[:3] == ["ACME Engineering", "Bridge Deck Assessment", "Prepared for: City of Springfield"]
    assert "Example Consulting" not in texts  # pandoc's title block was replaced by the cover

    used = {p.get(w("val")) for p in doc.iter(w("pStyle"))}
    assert "HouseBody" in used and "FirstParagraph" not in used and "BodyText" not in used
    assert {t.get(w("val")) for t in doc.iter(w("tblStyle"))} == {"LightGrid-Accent1"}

    # Every style the document uses is defined.
    defined = set(style_names(read(out, "word/styles.xml")))
    used_all = {e.get(w("val")) for tag in ("pStyle", "rStyle", "tblStyle") for e in doc.iter(w(tag))}
    assert used_all <= defined

    # The body section carries the template's main-section header.
    headers = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/tpl_header")]
    assert any(b"ACME Confidential" in zipfile.ZipFile(out).read(h) for h in headers)


def test_template_is_not_modified(files):
    tpl, md, out = files
    before = tpl.read_bytes()
    build(str(md), str(tpl), str(out))
    assert tpl.read_bytes() == before


def test_cli_learn_and_build(files, capsys, tmp_path):
    tpl, md, out = files
    assert main(["learn", str(tpl)]) == 0
    style_file = tmp_path / "style.json"
    style_file.write_text(capsys.readouterr().out)
    assert main(["build", str(md), "-t", str(tpl), "-s", str(style_file), "-o", str(out)]) == 0
    assert zipfile.is_zipfile(out)


def test_parts_without_a_content_type_are_registered(files, tmp_path):
    from rework.ooxml import Package, register_content_types

    tpl, _, _ = files
    pkg = Package(str(tpl))
    pkg.parts["word/fonts/House-Bold.ttf"] = b"font bytes"
    register_content_types(pkg, pkg)
    types = pkg.parts["[Content_Types].xml"].decode()
    assert 'Extension="ttf"' in types
    path = tmp_path / "fixed.docx"
    pkg.save(str(path))
    Document(str(path))  # python-docx refuses parts without a content type


def test_cover_may_use_a_heading_style(tmp_path):
    doc = Document()
    doc.add_heading("ACME LTD", level=1)  # some templates style the firm's name as a heading
    doc.add_paragraph("[Report name]")
    doc.add_page_break()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Body " * 200)
    path = tmp_path / "t.docx"
    doc.save(path)
    style = learn(str(path))
    assert style["cover"] == 3
    assert style["cover_lines"] == ["ACME LTD", "[Report name]"]


def test_long_front_matter_is_only_suggested(tmp_path):
    doc = Document()
    doc.add_paragraph("Statement " * 100)
    doc.add_page_break()
    doc.add_heading("Introduction", level=1)
    path = tmp_path / "t.docx"
    doc.save(path)
    style = learn(str(path))
    assert style["cover"] == 0 and "cover_note" in style


# Customer-test round 1: cover fields, headers, links and front matter.

FIELDS_MD = """---
title: Network Test
client: Northwind [Ltd]
date: 1 Oct 2026
abstract: Short summary of the work.
---

# Findings

Body text.
"""


def make_fields_template(path):
    """Placeholders written the way real templates write them: brackets, angle brackets, split runs."""
    doc = Document()
    doc.add_paragraph("Prepared for: [Client Name]")
    p = doc.add_paragraph("Date: ")
    p.add_run("[Report ").underline = True
    p.add_run("Date]").underline = True
    doc.add_paragraph("Delete this instruction")
    email = doc.add_paragraph()
    rid = email.part.relate_to("mailto:author@template.example", RT_HYPERLINK, is_external=True)
    link = etree.SubElement(email._p, w("hyperlink"), {f"{{{R_NS}}}id": rid})
    etree.SubElement(etree.SubElement(link, w("r")), w("t")).text = "<EMAIL>"
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "<client name> report"
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = "Confidential"
    doc.add_heading("Background", level=1)
    doc.add_paragraph("Body " * 50)
    doc.save(path)


@pytest.fixture
def fields(tmp_path):
    tpl, md = tmp_path / "fields.docx", tmp_path / "report.md"
    make_fields_template(tpl)
    md.write_text(FIELDS_MD)
    style = learn(str(tpl))
    style["cover"] = 5
    style["replace"] = [["[Client Name]", "{client}"], ["[Report Date]", "{date}"],
                        ["Delete this instruction", ""], ["<EMAIL>", "{email}"],
                        ["<client name>", "{client}"]]
    return tpl, md, tmp_path / "out.docx", style


def paragraphs(path, part="word/document.xml"):
    return ["".join(t.text or "" for t in p.iter(w("t"))) for p in read(path, part).iter(w("p"))]


def test_cover_fields_are_plain_text_and_come_from_front_matter(fields):
    tpl, md, out, style = fields
    warnings = build(str(md), str(tpl), str(out), style, {"email": "sam@firm.example"})
    texts = paragraphs(out)
    # "[Client Name]" is matched literally, and the inserted "[Ltd]" is not matched again.
    assert texts[0] == "Prepared for: Northwind [Ltd]"
    assert texts[1] == "Date: 1 Oct 2026"
    assert "Delete this instruction" not in texts and texts[2] == "sam@firm.example"  # emptied line dropped
    assert "Short summary of the work." in texts  # the abstract is kept
    assert warnings == []


def test_filling_keeps_the_placeholder_formatting(fields):
    tpl, md, out, style = fields
    build(str(md), str(tpl), str(out), style, {"email": "sam@firm.example"})
    date = list(read(out, "word/document.xml").iter(w("p")))[1]
    underlined = "".join(r.findtext(w("t")) or "" for r in date.iter(w("r")) if r.find(f".//{w('u')}") is not None)
    assert underlined == "1 Oct 2026"


def test_headers_are_filled_and_email_links_follow_the_text(fields):
    tpl, md, out, style = fields
    build(str(md), str(tpl), str(out), style, {"email": "sam@firm.example"})
    names = zipfile.ZipFile(out).namelist()
    headers = [n for n in names if "header" in n and n.endswith(".xml") and "_rels" not in n]
    assert headers and all("<client name>" not in "".join(paragraphs(out, h)) for h in headers)
    assert any("Northwind [Ltd] report" in paragraphs(out, h) for h in headers)
    rels = zipfile.ZipFile(out).read("word/_rels/document.xml.rels").decode()
    assert "mailto:sam@firm.example" in rels and "author@template.example" not in rels


def test_headers_come_before_footers_in_the_body_section(fields):
    tpl, md, out, style = fields
    build(str(md), str(tpl), str(out), style, {"email": "x@y.example"})
    sect = read(out, "word/document.xml").find(f"{w('body')}/{w('sectPr')}")
    tags = [etree.QName(e).localname for e in sect if "Reference" in etree.QName(e).localname]
    assert tags == sorted(tags, key=lambda t: t != "headerReference")


def test_missing_values_and_bad_pairs_are_errors(fields, capsys):
    tpl, md, out, style = fields
    with pytest.raises(ValueError, match="--set email="):
        build(str(md), str(tpl), str(out), style)
    style["replace"] = [{"pattern": "x", "replacement": "y"}]
    with pytest.raises(ValueError, match="list of"):
        build(str(md), str(tpl), str(out), style)
    with pytest.raises(SystemExit):  # argparse rejects --set without KEY=VALUE
        main(["build", str(md), "-t", str(tpl), "-o", str(out), "--set", "client"])


def test_text_that_is_not_found_is_reported(fields):
    tpl, md, out, style = fields
    style["replace"].append(["[Not There]", "x"])
    warnings = build(str(md), str(tpl), str(out), style, {"email": "x@y.example"})
    assert any("[Not There]" in warning for warning in warnings)


def test_style_names_work_in_style_json(files):
    tpl, md, out = files
    style = learn(str(tpl))
    style["body"] = "house body"  # a name, any case
    assert build(str(md), str(tpl), str(out), style) == []
    assert "HouseBody" in {p.get(w("val")) for p in read(out, "word/document.xml").iter(w("pStyle"))}


def test_body_section_is_where_the_chapters_are(tmp_path):
    doc = Document()
    doc.add_paragraph("Cover")
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.styles.add_style("Abstract Text", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("A long abstract. " * 300, style="Abstract Text")  # more text than the body
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Short.")
    doc.add_heading("Methods", level=1)
    path = tmp_path / "t.docx"
    doc.save(path)
    style = learn(str(path))
    assert style["body_section"] == 2
    assert style["body"] == "Normal"  # not the abstract's style


def test_unnumbered_headings_skip_the_template_numbering(files, tmp_path):
    tpl, _, out = files
    md = tmp_path / "u.md"
    md.write_text("# Abstract {.unnumbered}\n\nText.\n\n# Introduction\n\nSee [the abstract](#abstract).\n")
    build(str(md), str(tpl), str(out))
    doc = read(out, "word/document.xml")
    heads = {"".join(t.text or "" for t in p.iter(w("t"))): p for p in doc.iter(w("p"))
             if p.find(f"{w('pPr')}/{w('pStyle')}") is not None
             and p.find(f"{w('pPr')}/{w('pStyle')}").get(w("val")).startswith("Heading")}
    assert heads["Abstract"].find(f"{w('pPr')}/{w('numPr')}/{w('numId')}").get(w("val")) == "0"
    assert heads["Introduction"].find(f"{w('pPr')}/{w('numPr')}") is None
    names = {b.get(w("name")) for b in doc.iter(w("bookmarkStart"))}
    anchors = {h.get(w("anchor")) for h in doc.iter(w("hyperlink")) if h.get(w("anchor"))}
    assert anchors and anchors <= names  # the link to the heading still works


def test_captions_are_numbered_when_the_template_numbers_them(tmp_path):
    doc = Document()
    doc.add_heading("Results", level=1)
    caption = doc.add_paragraph("Table ")
    etree.SubElement(caption._p, w("fldSimple"), {w("instr"): " SEQ Table \\* ARABIC "})
    tpl, md, out = tmp_path / "t.docx", tmp_path / "r.md", tmp_path / "out.docx"
    doc.save(tpl)
    md.write_text("| a |\n|---|\n| 1 |\n\nTable: Culverts\n")
    style = learn(str(tpl))
    assert style["number_captions"] is True
    build(str(md), str(tpl), str(out), style, pandoc_args=["--toc"])
    assert b"SEQ Table" in zipfile.ZipFile(out).read("word/document.xml")


def test_images_are_found_next_to_the_source(files, tmp_path, monkeypatch):
    tpl, _, out = files
    folder = tmp_path / "notes"
    folder.mkdir()
    png = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                        "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")
    (folder / "pic.png").write_bytes(png)
    (folder / "r.md").write_text("![A figure](pic.png)\n")
    monkeypatch.chdir(tmp_path)  # build from another folder
    build("notes/r.md", str(tpl), str(out))
    assert any(n.startswith("word/media/") for n in zipfile.ZipFile(out).namelist())


def test_output_must_be_docx(files):
    tpl, md, _ = files
    with pytest.raises(SystemExit):
        main(["build", str(md), "-t", str(tpl), "-o", "report.dotx"])


def test_long_front_matter_lists_where_it_can_be_cut(tmp_path):
    doc = Document()
    doc.add_paragraph("Document Title")
    doc.add_paragraph("XX Months 20XX")
    for _ in range(5):
        doc.add_paragraph("Lorem ipsum " * 10)
    doc.add_page_break()
    doc.add_heading("Introduction", level=1)
    path = tmp_path / "t.docx"
    doc.save(path)
    style = learn(str(path))
    assert style["cover"] == 0 and style["cover_choices"]["2"] == "XX Months 20XX"


# Customer-test round 2.

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"


def test_the_readme_example_builds_without_warnings(tmp_path):
    style = json.loads((EXAMPLES / "acme-style.json").read_text())
    out = tmp_path / "acme.docx"
    warnings = build(str(EXAMPLES / "report.md"), str(EXAMPLES / "acme-template.docx"), str(out), style,
                     {"client": "Springfield County"})
    assert warnings == []
    assert "Prepared for Springfield County" in paragraphs(out)


def test_a_style_name_wins_over_another_styles_id():
    styles = etree.fromstring(
        f'<w:styles xmlns:w="{W_NS}">'
        '<w:style w:type="table" w:styleId="TableNormal"><w:name w:val="Normal Table"/></w:style>'
        '<w:style w:type="table" w:styleId="TableNormal0"><w:name w:val="TableNormal"/></w:style>'
        "</w:styles>")
    assert resolve_styles({"table": "TableNormal"}, styles)["table"] == "TableNormal0"


def test_the_body_starts_on_a_new_page_after_the_cover(files):
    tpl, md, out = files
    style = learn(str(tpl))
    style["cover"] = 3  # the three cover lines, without the section break that ends the page
    build(str(md), str(tpl), str(out), style)
    fourth = list(read(out, "word/document.xml").find(w("body")))[3]
    assert any(b.get(w("type")) == "page" for b in fourth.iter(w("br")))


def test_contents_heading_is_never_numbered(files):
    tpl, md, out = files
    build(str(md), str(tpl), str(out), pandoc_args=["--toc"])
    toc = [s for s in read(out, "word/styles.xml").iter(w("style")) if s.get(w("styleId")) == "TOCHeading"]
    assert toc and toc[0].find(f"{w('pPr')}/{w('numPr')}/{w('numId')}").get(w("val")) == "0"


def test_white_body_text_is_made_visible(tmp_path):
    doc = Document()
    doc.styles["Normal"].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # meant for a dark cover
    doc.add_heading("Section", level=1)
    doc.add_paragraph("Body " * 20)
    tpl, md, out = tmp_path / "t.docx", tmp_path / "r.md", tmp_path / "out.docx"
    doc.save(tpl)
    md.write_text("Plain text.\n")
    style = learn(str(tpl))
    style["body"] = "Normal"
    warnings = build(str(md), str(tpl), str(out), style)
    assert any("white" in warning for warning in warnings)
    colours = {c.get(w("val")) for c in read(out, "word/document.xml").iter(w("color"))}
    assert "auto" in colours


def test_links_are_kept_unless_a_replacement_covers_them(fields):
    tpl, md, out, style = fields
    warnings = build(str(md), str(tpl), str(out), style, {"email": "sam@firm.example"})
    assert not any("link" in warning for warning in warnings)  # a correct link raises nothing
    style["replace"] = [["Date: [Report Date]", "x"], ["<EMAIL>", ""]]
    warnings = build(str(md), str(tpl), str(out), style)
    assert any("link on '<EMAIL>' was removed" in warning for warning in warnings)


def test_captions_quotes_and_strict(tmp_path, capsys):
    doc = Document()
    doc.styles.add_style("Table or chart caption", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Table 1: sample", style="Table or chart caption")
    doc.add_paragraph("A quote.", style="Quote")
    tpl, md, out = tmp_path / "t.docx", tmp_path / "r.md", tmp_path / "out.docx"
    doc.save(tpl)
    md.write_text("| a |\n|---|\n| 1 |\n\nTable: Culverts\n\n> Quoted.\n")
    style = learn(str(tpl))
    assert style["table_caption"] == "Table or chart caption" and style["quote"] == "Quote"
    build(str(md), str(tpl), str(out), style)
    doc_xml = read(out, "word/document.xml")
    caption = next(p for p in doc_xml.iter(w("p")) if "Culverts" in "".join(t.text or "" for t in p.iter(w("t"))))
    assert caption.find(f"{w('pPr')}/{w('keepNext')}") is not None
    assert "Quote" in {p.get(w("val")) for p in doc_xml.iter(w("pStyle"))}
    style["replace"] = [["not in the template", "x"]]
    (tmp_path / "s.json").write_text(json.dumps(style))
    assert main(["build", str(md), "-t", str(tpl), "-s", str(tmp_path / "s.json"), "-o", str(out), "--strict"]) == 1


def test_learn_explains_bad_inputs(tmp_path, capsys):
    (tmp_path / "x.pdf").write_bytes(b"%PDF-1.4")
    assert main(["learn", str(tmp_path / "x.pdf")]) == 1
    assert "not a Word file" in capsys.readouterr().err
    assert main(["learn", str(tmp_path / "missing.docx")]) == 1


def test_cover_choices_show_breaks(tmp_path):
    doc = Document()
    doc.add_paragraph("Report title")
    doc.add_paragraph("Legal text " * 60)
    doc.add_page_break()
    doc.add_paragraph("Copyright 2026")
    doc.add_page_break()
    doc.add_heading("Introduction", level=1)
    path = tmp_path / "t.docx"
    doc.save(path)
    choices = learn(str(path))["cover_choices"]
    assert "[page break]" in choices["3"] and choices["4"] == "Copyright 2026"
